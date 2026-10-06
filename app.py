import os
import re
import io
import json
import base64
import tempfile
from typing import Optional, List, Dict, Any
import requests
from fastapi import FastAPI, Request, File, Form, UploadFile, HTTPException
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ---------------------------------------------------------
# FASTAPI UYGULAMA YAPILANDIRMASI
# ---------------------------------------------------------
app = FastAPI(
    title="HaberCiM - Ege Ajans AI Medya Asistanı",
    description="Geliştiren: Dr. Kemal ŞİMŞEK (Bilgisayar Mühendisi)",
    version="1.0.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Statik dosyalar ve şablonlar
os.makedirs("static", exist_ok=True)
os.makedirs("templates", exist_ok=True)

app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")

GEMINI_BASE_URL = "https://generativelanguage.googleapis.com/v1beta"
DEFAULT_RECTOR = "Prof. Dr. Musa ALCI"

FALLBACK_MODELS = [
    "gemini-2.5-flash",
    "gemini-3.8-flash",
    "gemini-1.5-flash",
    "gemini-2.0-flash",
    "gemini-2.5-pro",
]

# ---------------------------------------------------------
# YARDIMCI GEMINI API FONKSİYONLARI
# ---------------------------------------------------------
def call_gemini_api(api_key: str, primary_model: str, payload: dict, clean_stars: bool = False) -> dict:
    if not api_key:
        return {"success": False, "error": "API Anahtarı bulunamadı. Lütfen Ayarlar sekmesinden Gemini API anahtarınızı girin."}

    models_to_try = [primary_model] + [m for m in FALLBACK_MODELS if m != primary_model]
    last_error = ""

    for model in models_to_try:
        try:
            url = f"{GEMINI_BASE_URL}/models/{model}:generateContent?key={api_key.strip()}"
            response = requests.post(url, json=payload, headers={"Content-Type": "application/json"}, timeout=90)
            
            if response.status_code == 200:
                data = response.json()
                candidates = data.get("candidates", [])
                if candidates:
                    parts = candidates[0].get("content", {}).get("parts", [])
                    if parts:
                        text = parts[0].get("text", "")
                        if clean_stars:
                            text = text.replace("**", "").replace("*", "")
                        return {
                            "success": True,
                            "text": text.strip(),
                            "used_model": model
                        }
                last_error = "Modelden boş içerik döndü."
            else:
                try:
                    err_json = response.json()
                    msg = err_json.get("error", {}).get("message", f"HTTP {response.status_code}")
                    last_error = f"{model}: {msg}"
                except Exception:
                    last_error = f"{model}: HTTP {response.status_code} - {response.text[:200]}"
        except Exception as e:
            last_error = f"{model}: {str(e)}"

    return {"success": False, "error": f"Gemini API çağrısı başarısız oldu:\n{last_error}"}


def create_word_document(title: str, content: str, is_tv_format: bool = False) -> str:
    """Standart Microsoft Word (.docx) belgesi üretir."""
    doc = docx.Document()

    # Sayfa Kenar Boşlukları (1 inç)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Metaveri & İmza (Dr. Kemal ŞİMŞEK)
    core_props = doc.core_properties
    core_props.title = title
    core_props.author = "Dr. Kemal ŞİMŞEK - Bilgisayar Mühendisi"
    core_props.last_modified_by = "Dr. Kemal ŞİMŞEK - Bilgisayar Mühendisi"
    core_props.comments = "HaberCiM - Ege Üniversitesi Ege Ajans AI Asistanı ile üretilmiştir."

    # Normal Stil (Times New Roman 12 pt)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    # Başlık Ekleme (Eğer TV formatı değilse)
    if not is_tv_format and title:
        heading = doc.add_heading(level=1)
        run = heading.add_run(title.replace("#", "").strip())
        run.font.name = 'Times New Roman'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x00, 0x33, 0x66) # Ege Navy
        heading.paragraph_format.space_after = Pt(12)

    # Paragrafları İşleme
    for line in content.split('\n'):
        line_clean = line.strip()
        if not line_clean:
            continue

        if line_clean.startswith('# ') and not is_tv_format:
            h = doc.add_heading(level=1)
            r = h.add_run(line_clean[2:].strip())
            r.font.name = 'Times New Roman'
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
            h.paragraph_format.space_after = Pt(8)
        elif line_clean.startswith('## ') and not is_tv_format:
            h = doc.add_heading(level=2)
            r = h.add_run(line_clean[3:].strip())
            r.font.name = 'Times New Roman'
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x00, 0xA3, 0xE0)
            h.paragraph_format.space_after = Pt(6)
        elif line_clean.startswith('### ') and not is_tv_format:
            h = doc.add_heading(level=3)
            r = h.add_run(line_clean[4:].strip())
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.font.bold = True
            h.paragraph_format.space_after = Pt(4)
        elif line_clean.startswith('**') and line_clean.endsWith('**') and len(line_clean) > 4:
            p = doc.add_paragraph()
            r = p.add_run(line_clean[2:-2].strip())
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.font.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if not is_tv_format else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(8)
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if not is_tv_format else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(6)

            # Markdown **kalın** etiketlerini parse et
            parts = re.split(r'(\*\*.*?\*\*)', line_clean)
            for part in parts:
                if part.startswith('**') and part.endsWith('**') and len(part) > 4:
                    r = p.add_run(part[2:-2])
                    r.font.name = 'Times New Roman'
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                    r.font.name = 'Times New Roman'

    # Geçici dosyaya kaydet
    safe_title = re.sub(r'[\\/*?:"<>|]', '', title).replace(' ', '_').lower()[:35] or "habercim_belge"
    tmp_path = os.path.join(tempfile.gettempdir(), f"{safe_title}.docx")
    doc.save(tmp_path)
    return tmp_path

# ---------------------------------------------------------
# PYDANTIC İSTEK MODELLERİ
# ---------------------------------------------------------
class ApiKeyValidateRequest(BaseModel):
    api_key: str

class NewsGenerateRequest(BaseModel):
    api_key: str
    raw_text: str
    reference_url: Optional[str] = None
    rector_name: Optional[str] = DEFAULT_RECTOR
    model: Optional[str] = "gemini-2.5-flash"
    image_base64: Optional[str] = None
    image_mime_type: Optional[str] = None

class TvConvertRequest(BaseModel):
    api_key: str
    news_text: str
    has_video: bool = True
    model: Optional[str] = "gemini-2.5-flash"

class EditorialCheckRequest(BaseModel):
    api_key: str
    news_text: str
    rector_name: Optional[str] = DEFAULT_RECTOR
    model: Optional[str] = "gemini-2.5-pro"

class DocxExportRequest(BaseModel):
    title: str
    content: str
    is_tv: bool = False

# ---------------------------------------------------------
# API ENDPOINT'LERİ
# ---------------------------------------------------------

@app.get("/", response_class=HTMLResponse)
async def serve_home(request: Request):
    """Ana SPA Web Arayüzünü sunar."""
    index_path = os.path.join(os.path.dirname(__file__), "templates", "index.html")
    if os.path.exists(index_path):
        return FileResponse(index_path, media_type="text/html")
    return HTMLResponse("<h1>HaberCiM - Ege Ajans AI Platformu</h1>")

@app.get("/favicon.ico")
async def favicon():
    icon_path = os.path.join(os.path.dirname(__file__), "static", "logo.png")
    if os.path.exists(icon_path):
        return FileResponse(icon_path, media_type="image/png")
    return JSONResponse(status_code=204, content=None)

@app.get("/api/health")
async def health_check():
    return {
        "status": "online",
        "app": "HaberCiM - Ege Ajans AI Asistanı",
        "developer": "Dr. Kemal ŞİMŞEK - Bilgisayar Mühendisi",
        "version": "1.0.0"
    }

@app.post("/api/validate-key")
async def validate_api_key(req: ApiKeyValidateRequest):
    try:
        url = f"{GEMINI_BASE_URL}/models?key={req.api_key.strip()}"
        res = requests.get(url, timeout=10)
        if res.status_code == 200:
            return {"valid": True, "message": "API Anahtarı geçerli ve Google Gemini ile bağlantı kuruldu!"}
        return {"valid": False, "message": f"Geçersiz API Anahtarı (HTTP {res.status_code})"}
    except Exception as e:
        return {"valid": False, "message": f"Bağlantı hatası: {str(e)}"}

@app.post("/api/fetch-models")
async def fetch_remote_models(req: ApiKeyValidateRequest):
    try:
        url = f"{GEMINI_BASE_URL}/models?key={req.api_key.strip()}"
        res = requests.get(url, timeout=10)
        if res.status_code == 200:
            data = res.json()
            models = [
                m["name"].replace("models/", "")
                for m in data.get("models", [])
                if "gemini" in m.get("name", "")
            ]
            return {"success": True, "models": models}
        return {"success": False, "error": f"Modeller listelenemedi (HTTP {res.status_code})"}
    except Exception as e:
        return {"success": False, "error": str(e)}

@app.post("/api/transcribe-audio")
async def transcribe_audio(
    api_key: str = Form(...),
    model: str = Form("gemini-2.5-flash"),
    audio_file: Optional[UploadFile] = File(None),
    audio_base64: Optional[str] = Form(None),
    mime_type: Optional[str] = Form("audio/mp4")
):
    system_prompt = """Sen bir Türkçe Ses Çözümleme ve Deşifre Uzmanısın.
Görevin, sana verilen ses kaydını dikkatlice dinleyerek en yüksek doğrulukla Türkçe metne çevirmektir.
Kurallar:
1. Sesi tam ve doğru bir şekilde Türkçe metne dök.
2. Metindeki bariz dil bilgisi, harf ve yazım hatalarını düzelt, akıcılığı sağla.
3. SADECE nihai, iyileştirilmiş Türkçe metni ver.
4. Başka hiçbir açıklama, giriş, selamlama veya sonuç cümlesi ekleme.
5. Metni mantıksal paragraflara ayır."""

    raw_bytes = None
    final_mime = mime_type or "audio/mp4"

    if audio_file:
        raw_bytes = await audio_file.read()
        if audio_file.content_type:
            final_mime = audio_file.content_type
    elif audio_base64:
        if "," in audio_base64:
            audio_base64 = audio_base64.split(",")[1]
        raw_bytes = base64.b64decode(audio_base64)

    if not raw_bytes:
        raise HTTPException(status_code=400, detail="Ses verisi bulunamadı.")

    base64_data = base64.b64encode(raw_bytes).decode("utf-8")

    payload = {
        "system_instruction": {
            "parts": [{"text": system_prompt}]
        },
        "contents": [
            {
                "parts": [
                    {
                        "inline_data": {
                            "mime_type": final_mime,
                            "data": base64_data
                        }
                    },
                    {
                        "text": "Lütfen bu ses kaydını yukarıdaki kurallara göre harfiyen çözümle ve mantıksal paragraflar halinde yaz."
                    }
                ]
            }
        ],
        "generationConfig": {
            "temperature": 0.2
        }
    }

    result = call_gemini_api(api_key=api_key, primary_model=model, payload=payload)
    return result

@app.post("/api/generate-news")
async def generate_news(req: NewsGenerateRequest):
    rector = req.rector_name or DEFAULT_RECTOR
    prompt = f"""
Sen Ege Üniversitesi Ege Ajans (euegeajans.com) Haber Merkezi'nin kıdemli başyazarı ve haber editörüsün.
Görevin, sana sunulan ham notları, etkinlik bilgilerini ve (varsa) görseli kullanarak Ege Ajans'ın resmi ve saygın kurumsal yayın çizgisine tam uyumlu, kusursuz bir haber metni yazmaktır.

EGE AJANS KURUMSAL HABER FORMATI VE YAZIM STANDARTLARI:
1. Haber Mimarisi (Ters Piramit & 5N1K Kuralı):
   - Başlık: '# ' ile başlayan, büyük ve çarpıcı haber başlığı. (Örn: # Ege Üniversitesinden Sürdürülebilir Bilim Atağı)
   - Spot (Özet): '**' ile kalınlaştırılmış, haberin ana fikrini ve 5N1K unsurlarını özetleyen 1-2 cümlelik vurucu spot.
   - Mahreç ve Giriş: Haber gövdesinin ilk paragrafı MUTLAKA "**İZMİR (Ege Ajans) -** " ifadesiyle başlamalıdır. Giriş cümlesinde etkinliğin nerede, ne zaman, kimlerin katılımıyla gerçekleştiği net aktarılmalıdır.
   - Rektör / Yetkili Görüşü: Giriş paragrafını takiben, Ege Üniversitesi Rektörü {rector}'nın vizyoner açıklamalarına ve demeçlerine yer ver. Demeci tırnak içinde ("...") belirt ve cümlenin sonunu "... dedi", "... ifadelerini kullandı" veya "... şeklinde konuştu" ile bağla.
   - Ara Başlıklar: Konu akışını düzenleyen '## ' seviyesinde anlamlı ve kurumsal ara başlıklar kullan.
   - İçerik Vurguları: Ege Üniversitesi'nin araştırma üniversitesi misyonu, TÜBİTAK/uluslararası başarıları, öğrenci odaklı yaklaşımı, akreditasyonları ve topluma hizmet ilkelerini kurumsal dille vurgula.
   - Etiketler: En sonda '### Etiketler:' başlığı altında 4-6 adet hashtag ekle (Örn: #EgeÜniversitesi #EgeAjans #Bilim).
2. Dil ve Üslup:
   - Tarafsız, güvenilir, saygın, akıcı ve editoryal ajans dili.
   - Cümleler düşük olmamalı, TDK kurallarına tam uyumlu olmalıdır.
   - Asla modelin kendi sohbet ifadelerini ekleme. Yalnızca doğrudan haber metnini Markdown olarak üret.

Girdi Bilgileri:
Ham Metin / Etkinlik Notları: {req.raw_text}
Referans / Kaynak Bağlantı: {req.reference_url or "Belirtilmedi"}
"""
    parts = [{"text": prompt}]

    if req.image_base64 and req.image_mime_type:
        clean_b64 = req.image_base64.split(",")[1] if "," in req.image_base64 else req.image_base64
        parts.append({
            "inline_data": {
                "mime_type": req.image_mime_type,
                "data": clean_b64
            }
        })

    payload = {
        "contents": [{"parts": parts}],
        "generationConfig": {"temperature": 0.65}
    }

    result = call_gemini_api(api_key=req.api_key, primary_model=req.model or "gemini-2.5-flash", payload=payload)
    return result

@app.post("/api/convert-tv")
async def convert_tv(req: TvConvertRequest):
    if req.has_video:
        video_instruction = """
Haberde GÖRÜNTÜ / VTR / RÖPORTAJ MEVCUTTUR.
Biçimlendirme Kuralları:
1. KJ: Kısa, vurucu ve haberin özünü veren alt bant başlığı. (İlk satır 'KJ: ...' olmalıdır)
2. CAM: Stüdyo spikerinin kameraya bakarak okuyacağı, izleyiciyi habere bağlayan 1-2 cümlelik dinamik giriş.
3. SES: Görüntü (VTR) üzerine seslendirilecek 1-2 cümlelik dış ses metni.
4. Video / Röportaj İbreleri: Haberdeki önemli kişilerin konuşma geçişleri (Örn: Prof. Dr. ... Video).
5. SES: VTR devamı dış ses kapanış cümlesi.
"""
    else:
        video_instruction = """
Haberde GÖRÜNTÜ / VİDEO / VTR BULUNMAMAKTADIR. (SADECE STÜDYO SPİKERİ OKUYACAK)
Biçimlendirme Kuralları:
1. KJ: Kısa, dikkat çekici alt yazı başlığı. (İlk satır 'KJ: ...' olmalıdır)
2. CAM: Stüdyo spikerinin doğrudan kameraya bakarak haberi baştan sona özetleyeceği, akıcı, net ve vurucu bülten metni (3-4 kısa ve net cümle).
3. Kesinlikle VTR, Video, Röportaj, SES gibi var olmayan görüntü ibreleri EKLEME.
"""

    prompt = f"""
Sen usta bir TV haber editörü ve bülten koordinatörüsün. Sana verilen haber metnini televizyon haber bültenine uygun olarak **oldukça kısa, öz, vurucu ve net** bir şekilde yeniden yazmalısın.

{video_instruction}

ÖNEMLİ YAZIM KURALLARI:
- **ASLA markdown kalınlaştırma (**) veya yıldız işaretleri kullanma**. Örneğin CAM: veya KJ: gibi ifadelerde yıldız (**) asla kullanma.
- Çıktıyı tamamen temiz, düz metin (plain text) olarak üret.

Haber Metni:
---
{req.news_text}
---
"""
    payload = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {"temperature": 0.3}
    }

    result = call_gemini_api(api_key=req.api_key, primary_model=req.model or "gemini-2.5-flash", payload=payload, clean_stars=True)
    return result

@app.post("/api/check-editorial")
async def check_editorial(req: EditorialCheckRequest):
    rector = req.rector_name or DEFAULT_RECTOR
    prompt = f"""
Sen Ege Üniversitesi Ege Ajans Yayın Denetleme Kurulu Başkanı ve Baş Editörüsün.
Sana verilen haber metnini titizlikle incele ve 3 ana bölüme ayırarak değerlendir.

Aşağıdaki ETİKETLERİ VE BÖLÜM YAPISINI AYRI AYRI KULLANARAK YANIT VER:

=== RAPOR_BASLANGIC ===
# 📊 Ege Ajans Editoryal Değerlendirme Raporu

### 🎯 Genel Editoryal Puan: [100 üzerinden puan]/100
### 📋 5N1K Kontrol Karnesi
- **Ne?**: [Açıklama]
- **Kim?**: [Açıklama]
- **Nerede?**: [Açıklama]
- **Ne Zaman?**: [Açıklama]
- **Nasıl?**: [Açıklama]
- **Neden?**: [Açıklama]

### 🔍 TDK İmla ve Kurumsal Dil Denetimi
- İmla & Noktalama Durumu: [Değerlendirme]
- Ege Ajans Format Uyumu: [İzmir (Ege Ajans)- girişi, Rektör {rector} demeci ve tırnak alıntıları uyumu]
=== RAPOR_BITIS ===

=== DEGISIKLIKLER_BASLANGIC ===
### ✏️ Yapılan Düzeltmeler ve İyileştirme Maddeleri:
1. [Düzeltme 1]
2. [Düzeltme 2]
3. [Düzeltme 3]
=== DEGISIKLIKLER_BITIS ===

=== DUZELTILMIS_METIN_BASLANGIC ===
[Buraya haberin başlığı (# Başlık), spotu (**Spot**) ve Ege Ajans standartlarına göre sıfır hatayla yeniden yazılmış, kusursuz ve yayınlanmaya hazır nihai haber metnini yerleştir.]
=== DUZELTILMIS_METIN_BITIS ===

İncelenecek Haber Metni:
---
{req.news_text}
---
"""
    payload = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {"temperature": 0.2}
    }

    result = call_gemini_api(api_key=req.api_key, primary_model=req.model or "gemini-2.5-pro", payload=payload)
    
    if result.get("success"):
        raw_text = result.get("text", "")
        
        # Regex bölümleri ayır
        report_m = re.search(r"=== RAPOR_BASLANGIC ===([\s\S]*?)=== RAPOR_BITIS ===", raw_text)
        change_m = re.search(r"=== DEGISIKLIKLER_BASLANGIC ===([\s\S]*?)=== DEGISIKLIKLER_BITIS ===", raw_text)
        edited_m = re.search(r"=== DUZELTILMIS_METIN_BASLANGIC ===([\s\S]*?)=== DUZELTILMIS_METIN_BITIS ===", raw_text)

        report = report_m.group(1).strip() if report_m else raw_text
        changelog = change_m.group(1).strip() if change_m else ""
        edited_text = edited_m.group(1).strip() if edited_m else raw_text

        result["structured"] = {
            "report": report,
            "changelog": changelog,
            "edited_text": edited_text
        }

    return result

@app.post("/api/export-docx")
async def export_docx(req: DocxExportRequest):
    try:
        file_path = create_word_document(title=req.title, content=req.content, is_tv_format=req.is_tv)
        filename = os.path.basename(file_path)
        return FileResponse(
            path=file_path,
            filename=filename,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Word belgesi oluşturulamadı: {str(e)}")

# ---------------------------------------------------------
# ÇALIŞTIRMA (Yerel geliştirme için)
# ---------------------------------------------------------
if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run("app:app", host="0.0.0.0", port=port, reload=True)
